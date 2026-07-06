"""
Module Information Page Creator Script
@author: nnrw

This script automates the creation of HTML pages about Module Information. The HTML can be manually dropped into Canvas Pages if desired. 
"""
from dotenv import load_dotenv
import os
import sys
import re
import logging
from typing import Dict, Union
import openpyxl
from datetime import datetime
from openpyxl import load_workbook

# --- 1. CONFIGURATION & ENVIRONMENT SETUP ---
load_dotenv('File Addresses.env')
XLSX_FOLDER = os.getenv('XLSX_MOF_FILES_DIR')
AGGREGATED_MODULE_DATA = os.getenv("AGGREGATED_MODULE_DATA")
OUTPUT_HTML_FOLDER = os.getenv("MODULE_INFO_OUTPUT_DIR")

def log_print(message: str, file=sys.stdout):
    """
    Standardizes printing to the console for easier logging/debugging.
    """
    print(message, file=file)

def safe_read_cell(sheet: openpyxl.worksheet.worksheet.Worksheet, 
                   row_idx: int, 
                   col_idx: int) -> Union[str, int, float, None]:
    """
    Safely retrieves a cell value from an openpyxl sheet, stripping whitespace 
    from strings and handling empty cells gracefully.
    """
    try:
        cell = sheet.cell(row=row_idx, column=col_idx)
        value = cell.value
        if value is None:
            return None
        if isinstance(value, str):
            stripped = value.strip()
            return stripped if stripped else None
        return value
    except Exception:
        return None

def extract_module_content(module_code: str) -> Dict[str, Dict[str, str]]:
    """
    Opens the specific MOF Excel file for a module and searches for syllabus 
    sections (Aims, Outcomes, etc.). Returns a dictionary of formatted HTML content.
    """
    log_print(f"Retrieving information about {module_code} module aims, syllabus, knowledge and skills outcomes directly from XLSX version of MOF...")
    log_print("")
    file_path = os.path.join(XLSX_FOLDER, f"{module_code}.xlsx")
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Source XLSX not found: {file_path}")


    SEARCH_CONFIG = [
        ("aims", "Aims", "🎯", 11),  
        ("outline of syllabus", "Syllabus Outline", "📖", 13),
        ("intended knowledge outcomes", "Intended Knowledge Outcomes", "🎓", 11),
        ("intended skills outcomes", "Intended Skills Outcomes", "💡", 12),
    ]
    content_map: Dict[str, Dict[str, str]] = {} 
    try:
        wb = load_workbook(file_path) 
        ws = wb.active
        search_keywords = {term[0] for term in SEARCH_CONFIG}
        for row in ws.iter_rows():
            row_values = [c.value for c in row]
            for cell_index, cell in enumerate(row):
                if isinstance(cell.value, str):
                    # Clean the cell content to find keyword matches
                    search_value = re.sub(r'[\r\n]', ' ', cell.value).strip().lower()
                    if search_value in search_keywords:
                        for keyword, base_title, icon, offset in SEARCH_CONFIG:
                            if search_value == keyword and base_title not in content_map:
                                found_content_value = next((row_values[i] for i in range(cell_index + 1, len(row_values)) if row_values[i]), None)
                                content = str(found_content_value).replace('\n', '<br>') if found_content_value else "Content not found."
                                content_map[base_title] = {'icon': icon, 'content': content}
                                status_icon = "✅" if found_content_value else "⚠️"
                                summary = content[:50].replace('<br>', ' ') + "..." if found_content_value else "Empty"
                                log_print(f"    {status_icon} {icon} {base_title}: {summary}")
                                break
        # Add placeholders for any missing sections
        for keyword, base_title, icon, offset in SEARCH_CONFIG:
            if base_title not in content_map:
                content_map[base_title] = {'icon': icon, 'content': "<strong>Section not found in source.</strong>"}
    except Exception as e:
        log_print("")
        log_print(f"❌ ERROR reading XLSX for {module_code}: {e}")
        return {term[1]: {'icon': term[2], 'content': "Error reading file."} for term in SEARCH_CONFIG}
    return content_map

def generate_page_html(module_code: str, content_data: Dict[str, Dict[str, str]]) -> str:
    """
    Wraps the extracted module content in standard HTML formatting for Canvas,
    including a real-time generation timestamp.
    """
    # 1. Generate the missing timestamp string
    current_time_str = datetime.now().strftime("%d %B %Y at %H:%M")
    
    html_sections = []
    section_titles = ["Aims", "Syllabus Outline", "Intended Knowledge Outcomes", "Intended Skills Outcomes"]
    
    # 2. Build the main content blocks
    for title in section_titles:
        data = content_data.get(title)
        if data:
            html_sections.append(
                f'<span id="HTMLAnchor1" style="font-size: 12pt;"><h2><span style="font-size: 18pt;"><strong>{data["icon"]} {module_code} {title}</strong></span></h2>'
                f'<p>{data["content"]}</p><span id="HTMLAnchor2" style="font-size: 12pt;">'
            )
            html_sections.append('<hr style="margin-top: 30px;">')
    
    if html_sections: 
        html_sections.pop() # Remove the very last horizontal rule for a cleaner finish
        
    main_body = "".join(html_sections)
    admin_contact = os.getenv("ADMIN_CONTACT")
    
    return f"""{main_body}    
    <hr style="margin-top: 20px;">
    <p style="margin-top: 20px; font-size: small; color: #6b7280;">This content was generated for module code: <strong>{module_code}</strong> on {current_time_str}. If you notice any errors or can help us improve the accuracy of the information provided, please contact {admin_contact} to let us know.</p>
    """

def get_module_codes_from_reference_folder(folder_path):
    """
    Returns a list of module codes derived from filenames in the reference folder.
    Expected filename format: {module_code}.xlsx
    """
    module_codes = []
    for filename in os.listdir(folder_path):
        if filename.lower().endswith(".xlsx"):
            module_code = os.path.splitext(filename)[0]
            module_codes.append(module_code)
    return sorted(set(module_codes))

if __name__ == "__main__":

    page_urls_for_excel = []
    logging.getLogger('canvasapi').setLevel(logging.WARNING)
    logging.getLogger('urllib3').setLevel(logging.WARNING)

    try:
        os.makedirs(OUTPUT_HTML_FOLDER, exist_ok=True)

        module_codes = get_module_codes_from_reference_folder(XLSX_FOLDER)
        log_print(f"📁 Found {len(module_codes)} module reference files")
        log_print(f"📂 XLSX folder → {XLSX_FOLDER}")
        log_print(f"📂 HTML output → {OUTPUT_HTML_FOLDER}")

        pages_processed = 0
        pages_successful = 0

        for i, module_code in enumerate(module_codes, start=1):
            try:
                log_print(f"▶ [{i}/{len(module_codes)}] Processing {module_code}")
                content_data = extract_module_content(module_code)
                HTML_CONTENT = generate_page_html(
                    module_code,
                    content_data,
                )

                if not HTML_CONTENT.strip():
                    raise ValueError("Empty HTML content generated")
                pages_processed += 1
                safe_module_code = module_code.replace(" ", "_")
                output_path = os.path.join(
                    OUTPUT_HTML_FOLDER,
                    f"{safe_module_code}_Module_Information.html"
                )

                with open(output_path, "w", encoding="utf-8") as f:
                    f.write(HTML_CONTENT)

                pages_successful += 1
                log_print(f"✅ Written → {output_path}")

            except Exception as e:
                log_print(f"❌ Module {module_code} failed: {e}")

    except KeyboardInterrupt:
        log_print("\n🛑 Stopped by user.")
    finally:
        log_print("\n--- Summary ---")
        log_print(f"Processed: {pages_processed}")
        log_print(f"Successful: {pages_successful}")
        log_print("✅ Done.")
